
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
plt.rcParams["font.family"] = "Noto Sans CJK JP"
import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
INPUT_DIR = BASE_DIR / "input"
OUTPUT_DIR = BASE_DIR / "output"
REQUIRED_COLUMNS = ["日付", "商品名", "数量", "単価"]


def main():
    if not INPUT_DIR.is_dir():
        print(f"エラー: 入力フォルダがありません: {INPUT_DIR}")
        return

    files = sorted(
        p for p in INPUT_DIR.glob("*.xlsx")
        if not p.name.startswith("~$")
    )
    if not files:
        print("エラー: input フォルダにExcelファイルがありません。")
        return

    errors = []
    frames = []
    read_rows = valid_rows = excluded_rows = 0
    successful_files = excluded_files = 0

    for path in files:
        try:
            df = pd.read_excel(path, dtype=object)
        except Exception as exc:
            excluded_files += 1
            errors.append(f"{path.name}: 読み込み失敗: {exc}")
            continue

        read_rows += len(df)
        df.columns = [str(c).strip() for c in df.columns]
        missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]

        if missing:
            excluded_files += 1
            excluded_rows += len(df)
            errors.append(f"{path.name}: 必須列不足: {', '.join(missing)}")
            continue

        df = df[REQUIRED_COLUMNS].copy()
        df["商品名"] = df["商品名"].map(
            lambda x: x.strip() if isinstance(x, str) else x
        )

        dates = pd.to_datetime(df["日付"], errors="coerce")
        quantities = pd.to_numeric(df["数量"], errors="coerce")
        prices = pd.to_numeric(df["単価"], errors="coerce")
        product_ok = df["商品名"].notna() & (
            df["商品名"].astype(str).str.strip() != ""
        )
        valid = dates.notna() & quantities.notna() & prices.notna() & product_ok

        for i in df.index[~valid]:
            reasons = []
            if pd.isna(dates.loc[i]):
                reasons.append("日付が不正")
            if pd.isna(quantities.loc[i]):
                reasons.append("数量が不正")
            if pd.isna(prices.loc[i]):
                reasons.append("単価が不正")
            if not product_ok.loc[i]:
                reasons.append("商品名が空欄")
            errors.append(
                f"{path.name}: Excel行 {i + 2}: {', '.join(reasons)}"
            )

        excluded_rows += int((~valid).sum())
        clean = pd.DataFrame({
            "日付": dates[valid],
            "商品名": df.loc[valid, "商品名"],
            "数量": quantities[valid],
            "単価": prices[valid],
        })
        clean["売上金額"] = clean["数量"] * clean["単価"]
        clean["元ファイル"] = path.name

        if not clean.empty:
            frames.append(clean)
            valid_rows += len(clean)
            successful_files += 1
        else:
            excluded_files += 1
            errors.append(f"{path.name}: 有効なデータ行がありません。")

    if not frames:
        write_process_report(
            len(files), successful_files, excluded_files,
            read_rows, valid_rows, excluded_rows, errors
        )
        print("エラー: 有効な売上データがありません。")
        return

    sales = pd.concat(frames, ignore_index=True)
    sales["月"] = sales["日付"].dt.strftime("%Y-%m")

    monthly = (
        sales.groupby("月", as_index=False)["売上金額"]
        .sum().sort_values("月")
    )
    products = (
        sales.groupby("商品名", as_index=False)["売上金額"]
        .sum().sort_values("売上金額", ascending=False)
    )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    monthly_png = OUTPUT_DIR / "monthly_sales.png"
    products_png = OUTPUT_DIR / "product_sales.png"

    plt.figure(figsize=(5, 3))
    plt.bar(monthly["月"], monthly["売上金額"])
    plt.title("月別売上")
    plt.xlabel("月")
    plt.ylabel("売上金額（円）")
    plt.tight_layout()
    plt.savefig(monthly_png, dpi=150)
    plt.close()

    plt.figure(figsize=(5, 3))
    plt.bar(products["商品名"], products["売上金額"])
    plt.title("商品別売上")
    plt.xlabel("商品名")
    plt.ylabel("売上金額（円）")
    plt.tight_layout()
    plt.savefig(products_png, dpi=150)
    plt.close()

    report_path = OUTPUT_DIR / "sales_report.xlsx"
    with pd.ExcelWriter(report_path, engine="openpyxl") as writer:
        sales.to_excel(writer, sheet_name="全売上データ", index=False)
        monthly.to_excel(writer, sheet_name="月別売上", index=False)
        products.to_excel(writer, sheet_name="商品別売上", index=False)

    wb = load_workbook(report_path)
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for cells in sheet.columns:
            width = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in cells
            )
            sheet.column_dimensions[
                get_column_letter(cells[0].column)
            ].width = min(width + 2, 30)
    for row in wb["全売上データ"].iter_rows(min_row=2):
        row[0].number_format = "yyyy/mm/dd"
        row[2].number_format = '#,##0'
        row[3].number_format = '"¥"#,##0'
        row[4].number_format = '"¥"#,##0'

    for sheet_name in ["月別売上", "商品別売上"]:
        sheet = wb[sheet_name]
        for row in sheet.iter_rows(min_row=2):
            row[1].number_format = '"¥"#,##0'


    for sheet_name in ["月別売上", "商品別売上"]:
        sheet = wb[sheet_name]
        for row in sheet.iter_rows(min_row=2):
            row[1].number_format = '¥#,##0'

    wb["月別売上"].add_image(ExcelImage(str(monthly_png)), "D2")
    wb["商品別売上"].add_image(ExcelImage(str(products_png)), "D2")
    wb["全売上データ"].column_dimensions["B"].width = 12
    wb["全売上データ"].column_dimensions["D"].width = 12
    wb["全売上データ"].column_dimensions["E"].width = 14
    wb["月別売上"].column_dimensions["B"].width = 16
    wb.save(report_path)

    write_process_report(
        len(files), successful_files, excluded_files,
        read_rows, valid_rows, excluded_rows, errors
    )
    print("\n処理が完了しました。")
    print(f"Excelレポート: {report_path}")
    print(f"処理レポート: {OUTPUT_DIR / 'process_report.txt'}")


def write_process_report(
    total, successful, excluded, read_rows, valid_rows, excluded_rows, errors
):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    lines = [
        "Excel売上データ自動集計・レポート作成ツール 処理レポート",
        f"読み込んだファイル数: {total}",
        f"正常に処理したファイル数: {successful}",
        f"除外したファイル数: {excluded}",
        f"読み込んだ行数: {read_rows}",
        f"有効データ行数: {valid_rows}",
        f"除外した行数: {excluded_rows}",
        "",
        "エラー内容:",
    ]
    lines.extend(errors or ["エラーはありません。"])
    (OUTPUT_DIR / "process_report.txt").write_text(
        "\n".join(lines) + "\n", encoding="utf-8"
    )
    print("\n".join(lines))


if __name__ == "__main__":
    main()